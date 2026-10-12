#!/usr/bin/env python3
"""
Convert the Lockbox "Manager Report Detail" CSV exports to Parquet.

Handles three source folders, each with its own column pruning and value
cleanup, ported from the Power Query that used to read it:

    Lockbox 12month raw data   -> monthly snapshots, feeds the MONTHLY report
    Current Raw Data           -> feeds the WEEKLY report
    All Data                   -> full history since 2020 (no derived columns;
                                  converted file by file)

Also rebuilds "Control Sheet data (parquet)" from the daily UMB control sheet
workbooks (see the CONTROL SHEET section).

One .parquet per source .csv, so a monthly rotation only costs you the new
file. Re-running is cheap: a file is rebuilt only when its CSV is newer than
the .parquet beside it.

After converting, it recomputes the current month's row in KPI Benchmarks.xlsx
(see the KPI BENCHMARKS section) and prints old vs new values.

Requires: pandas, pyarrow, openpyxl
    pip install pandas pyarrow openpyxl

Usage:
    python lockbox_to_parquet.py                 # convert everything
    python lockbox_to_parquet.py --audit         # compare CSV schemas, then exit
    python lockbox_to_parquet.py --force         # rebuild regardless of timestamps
    python lockbox_to_parquet.py --only weekly   # one profile (monthly/weekly/alldata)
    python lockbox_to_parquet.py --benchmarks-only --dry-run   # calibrate, write nothing
    python lockbox_to_parquet.py --no-benchmarks                # Parquet only
    python lockbox_to_parquet.py --only controlsheet            # control sheet step only
    python lockbox_to_parquet.py --calib-months 2026-06         # calibrate on one row
    python lockbox_to_parquet.py --force-benchmarks             # write even on a poor fit
"""

from __future__ import annotations

import argparse
import datetime
import re
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(r"Q:\CORE\PURCHASE TRADING\MANAGEMENT REPORTS (LOCKBOX)")
HOLIDAYS_CSV = ROOT / "holidays.csv"

# ---------------------------------------------------------------------------
# BUG_COMPAT_MODE
#
# False (default) = two bugs in the original DAX are FIXED:
#
#   1. DATEDIFF against a blank previous-timestamp returned ~1,104,586 hours
#      instead of blank, because DAX reads BLANK() as 1899-12-30. It fired on
#      every user's first batch in the window and dragged any average over the
#      hour-fraction columns. Now blank, which AVERAGE correctly skips.
#      Watch: COUNT/COUNTROWS over those columns will return fewer rows.
#
#   2. Hold Duration's NOT ISBLANK(StartDate) test was dead code -- StartDate
#      is INITIALPROBLEMDATE + 1, and blank + 1 = 1 in DAX, never blank. A row
#      with INITIALPROBLEMTYPE = "Hold" and a null INITIALPROBLEMDATE scored
#      45,969 days, which then tripped the Hold Duration >= 4 branch in Hold
#      Reason Categories and mislabelled it "Stop Pay Pending". Now 0.
#
# True = reproduce the old output exactly. Set it, rebuild, and diff visuals
# against the current model to confirm nothing else moved; then set it back.
#
# NOT a bug, deliberately unchanged: Employer_Earlist_Remit_timestamp groups by
# BATCHID despite its name. Batch is the remittance event and EMPLOYERNAME is
# the grouping applied by the next column, so the pair is correct as written.
# ---------------------------------------------------------------------------
BUG_COMPAT_MODE = False

# Bump whenever the derived-column logic changes. Each output folder records the
# version it was built with; a mismatch rebuilds that folder on the next run, so
# a logic change never needs --force. (2: case-insensitive matching like DAX,
# and hold-reason relabels. 3: no-reason holds and "NONE" labelled "Misc".
# 4: ...spelled out as "Miscellaneous".)
DERIVE_VERSION = 4

DAX_BLANK_DATE = pd.Timestamp("1899-12-30")


# ---------------------------------------------------------------------------
# building blocks
# ---------------------------------------------------------------------------

def group_min(df: pd.DataFrame, key: str, value: str) -> pd.Series:
    """CALCULATE(MIN([value]), ALLEXCEPT(Table, [key]))"""
    return df.groupby(key, dropna=False)[value].transform("min")


def previous_distinct_in_group(df: pd.DataFrame, key: str, value: str) -> pd.Series:
    """CALCULATE(MAX([value]), FILTER(Table, [key] = EARLIER([key])
                                      && [value] < EARLIER([value])))

    For each row: the largest [value] within its [key] group that is strictly
    less than this row's [value]. MAX over strictly-smaller values is the
    previous *distinct* value, so duplicates (rows sharing a BATCHID, and so
    sharing a timestamp) must all receive the same answer -- a row-wise
    shift() would be wrong on every tie.

    Rows with a null [value] get null, matching DAX: comparing against a blank
    EARLIER() value makes the FILTER empty, so MAX returns BLANK().
    """
    # DAX treats blank = blank as TRUE, so null keys form their own group.
    sentinel = "\x00__null__"
    keys = df[key].astype("object").where(df[key].notna(), sentinel)

    pairs = pd.DataFrame({key: keys, value: df[value]})
    pairs = pairs.dropna(subset=[value]).drop_duplicates()
    pairs = pairs.sort_values([key, value], kind="mergesort")
    pairs["__prev"] = pairs.groupby(key, sort=False, dropna=False)[value].shift(1)

    lookup = pd.DataFrame({key: keys, value: df[value]})
    merged = lookup.merge(pairs, on=[key, value], how="left")
    out = merged["__prev"]
    out.index = df.index
    return out


def hour_fraction_diff(start: pd.Series, end: pd.Series) -> pd.Series:
    """ABS(DATEDIFF([start], [end], SECOND)) / 3600"""
    if BUG_COMPAT_MODE:
        s = start.fillna(DAX_BLANK_DATE)
        e = end.fillna(DAX_BLANK_DATE)
    else:
        s, e = start, end
    return (e - s).dt.total_seconds().abs() / 3600


def datediff_days(start: pd.Series, end: pd.Series) -> pd.Series:
    """DATEDIFF([start], [end], DAY) -- counts day boundaries crossed, so the
    time component is dropped before subtracting."""
    return (end.dt.normalize() - start.dt.normalize()).dt.days


def format_hour_ampm(ts: pd.Series) -> pd.Series:
    """FORMAT(TIME(HOUR([ts]),0,0), "h AM/PM")  ->  "9 AM", "12 PM", "1 PM"

    HOUR(BLANK()) = 0 in DAX, so nulls come out as "12 AM".
    """
    hours = ts.dt.hour.fillna(0).astype(int)
    hour12 = hours % 12
    hour12 = hour12.where(hour12 != 0, 12)
    suffix = np.where(hours < 12, "AM", "PM")
    return pd.Series(
        hour12.astype(str).to_numpy() + " " + suffix, index=ts.index, dtype="string"
    )


def combined_status(status: pd.Series, issue: pd.Series) -> pd.Series:
    """CombinedStatus&InitialProb

    LEN(BLANK()) = 0 in DAX, so nulls count as blank.
    """
    s = status.fillna("").astype("string").str.strip()
    i = issue.fillna("").astype("string").str.strip()
    s_blank = s.str.len() == 0
    i_blank = i.str.len() == 0

    out = pd.Series(pd.NA, index=status.index, dtype="string")
    out = out.mask(~s_blank & ~i_blank, s + " Problem: " + i)
    out = out.mask(~s_blank & i_blank, s)
    out = out.mask(s_blank & ~i_blank, "Problem: " + i)
    # both blank stays <NA>, matching SWITCH -> BLANK()
    return out


def _ci(s: pd.Series) -> pd.Series:
    """Lower-cased text for comparisons. DAX compares text case-INSENSITIVELY
    ("NIGO" = "Nigo" is TRUE), so every ported text comparison goes through
    this. Spaces still count, as they do in DAX."""
    return s.fillna("").astype("string").str.lower()


def _ci_list(values) -> list:
    return [v.lower() for v in values]


def hold_duration(
    combined: pd.Series,
    initial_problem_date: pd.Series,
    processing_date: pd.Series,
    holidays: np.ndarray,
) -> pd.Series:
    """Hold Duration

    The original's NOT ISBLANK(StartDate) test was dead code (blank + 1 = 1 in
    DAX, never blank), so a null INITIALPROBLEMDATE scored ~45,969 days. Unless
    BUG_COMPAT_MODE is on, a null start now yields 0 as the author intended.
    """
    is_hold = _ci(combined).isin(_ci_list(["Processed Problem: Hold", "Reimburse Lockbox Problem: Hold"]))

    start = initial_problem_date.fillna(DAX_BLANK_DATE) + pd.Timedelta(days=1)
    end = processing_date

    total_days = datediff_days(start, end)

    # COUNTROWS('Holiday Table') filtered to [Date] >= Start && [Date] <= End.
    # Holiday dates sit at midnight; StartDate keeps its time component, so
    # compare raw timestamps as DAX does.
    hol = np.sort(np.asarray(holidays, dtype="datetime64[ns]"))
    s_vals = start.to_numpy(dtype="datetime64[ns]")
    e_vals = end.to_numpy(dtype="datetime64[ns]")
    lo = np.searchsorted(hol, s_vals, side="left")
    hi = np.searchsorted(hol, e_vals, side="right")
    holidays_between = pd.Series(
        np.where(pd.isna(e_vals), 0, np.maximum(hi - lo, 0)), index=combined.index
    )

    net = (total_days - holidays_between).clip(lower=0)
    keep = is_hold & end.notna() & (end >= start)
    if not BUG_COMPAT_MODE:
        keep = keep & initial_problem_date.notna()
    return net.where(keep, 0).astype("Int64")


def days_until_return(
    status: pd.Series, return_complete: pd.Series, processing_date: pd.Series
) -> pd.Series:
    """DaysuntilRETURN"""
    eligible = _ci(status) == "return complete - money returned"
    diff = datediff_days(return_complete, processing_date).abs()
    return diff.where(eligible).astype("Float64")


def weekday_name(d: pd.Series) -> pd.Series:
    """FORMAT([PROCESSINGDATE], "dddd")"""
    return d.dt.day_name().astype("string")


# ---------------------------------------------------------------------------
# Hold Reason Categories
#
# Branch order matters -- SWITCH returns the first match. HOLD REASON = "" and
# INITIALPROBLEMTYPE = "" also match nulls, per DAX blank/"" equivalence.
# ---------------------------------------------------------------------------

_HR_DISCLOSURE = [
    "Nigo", "Not Disclosed", "Received less than expected",
    "Received more than expected", "Returned to Rep", "Rollover Not Set up",
    "Suit Review", "Suitability Review", "Which Account?", "Which RO",
    "Which Rollover?",
]
_HR_EMPLOYER = [
    "NIGO - New Employer", "No Account Established", "No Payslot",
    "Remittance Doesn't Match Check/ACH", "Remittance Unclear/Missing",
    "Wrong SSN",
]
_HR_FUNDS_RETURNED = ["Check Returned to PM", "Refund of Contribution", "Symetra $30 Min"]
_HR_IT = ["Batch Not Allocating", "Invalid Product", "IT Issue", "No ER attached to Account"]
_HR_NO_ACCOUNT = ["Account Auto Closed", "Not in OES/CB", "Not in OES or ClientBridge"]
_HR_ACCOUNT_SETUP = ["Deferral Not Set Up", "Waiting for NA"]


# Labels for the two kinds of "no category" the original SWITCH produced.
#   * A hold with no HOLD REASON entered. The original labelled these
#     "Miscellaneous" (its comments say it meant "uncategorized"), which made
#     Miscellaneous by far the biggest category while almost no row actually
#     says Miscellaneous or NONE.
#   * A HOLD REASON that isn't in any list below. The original had no else, so
#     these showed as a blank category; now they're named so a new reason
#     someone starts using shows up instead of disappearing.
HR_LABEL_NO_REASON = "Miscellaneous"        # per manager (was briefly "No Hold Reason Entered", then "Misc")
HR_LABEL_UNLISTED = "Unlisted Reason"


def hold_reason_categories(
    hold_reason: pd.Series,
    initial_problem_type: pd.Series,
    hold_dur: pd.Series,
) -> pd.Series:
    hr = _ci(hold_reason)
    ipt = _ci(initial_problem_type)
    hd = hold_dur.fillna(0)

    out = pd.Series(pd.NA, index=hold_reason.index, dtype="string")

    def branch(mask, label):
        nonlocal out
        out = out.where(out.notna() | ~mask, label)

    branch(hr.isin(_ci_list(_HR_DISCLOSURE)), "Disclosure Issue")
    branch(hr.isin(_ci_list(_HR_EMPLOYER)), "Employer Issue")
    branch(hr.isin(_ci_list(_HR_FUNDS_RETURNED)), "Funds Returned to PM")
    branch(hr.isin(_ci_list(_HR_IT)), "IT Issue")
    branch(hr == "loan", "Loan")
    branch(hr.isin(_ci_list(_HR_NO_ACCOUNT)), "No Account")
    branch(hr == "reimburse lockbox", "Reimburse Lockbox")
    branch(hr == "none", HR_LABEL_NO_REASON)   # an explicit "NONE" means the same thing
    branch(hr.isin(_ci_list(_HR_ACCOUNT_SETUP)), "Account Setup Issue")
    branch(hr == "tpa approval needed", "TPA Approval Needed")
    # && binds tighter than ||. Note the original's own rule: a hold with no
    # reason lasting 4+ days is treated as a stop pay.
    branch(
        (hr == "stop pay pending") | ((hr == "") & (ipt == "hold") & (hd >= 4)),
        "Stop Pay Pending",
    )
    branch(hr == "umb deposit error", "UMB Deposit Error")
    branch((hr == "") & (ipt == ""), "No Holds on Batch")
    branch((hr == "") & (ipt == "hold"), HR_LABEL_NO_REASON)
    branch(ipt == "return", "Batch Returned")
    branch(hr != "", HR_LABEL_UNLISTED)
    return out


def user_group(user: pd.Series, members: list, label: str) -> pd.Series:
    """CREATEDBYUSER (groups) / Auditors

    Note the two originals disagree: the label is singular in one and plural in
    the other, and the member lists differ. Both preserved as written.
    """
    u = user.astype("string")
    out = u.copy()
    out = out.mask(_ci(u).isin(_ci_list(members)), label)
    out = out.mask(u.isna(), "(Blank)")
    return out.astype("string")


# ---------------------------------------------------------------------------
# holidays
# ---------------------------------------------------------------------------

def load_holidays(path) -> np.ndarray:
    df = pd.read_csv(path, encoding="utf-8-sig")
    col = "Date" if "Date" in df.columns else df.columns[0]
    return pd.to_datetime(df[col]).to_numpy(dtype="datetime64[ns]")


# ---------------------------------------------------------------------------
# per-profile assembly
#
# Column names differ between the two tables for identical logic, so each
# profile supplies its own names. Order matters: later columns read earlier ones.
# ---------------------------------------------------------------------------

def derive_monthly(df: pd.DataFrame, holidays: np.ndarray) -> pd.DataFrame:
    df["EarliestBatchPROCESSINGtimestamp"] = group_min(df, "BATCHID", "CREATEDDATE")
    df["HR_BatchedProcessed"] = format_hour_ampm(df["CREATEDDATE"])
    df["PreviousBatch_ProcessingTime"] = previous_distinct_in_group(
        df, "CREATEDBYUSER", "EarliestBatchPROCESSINGtimestamp"
    )
    df["ADUIT_Earliesttimestamp"] = group_min(df, "BATCHID", "AUDITDATE")
    df["PreviousAUDIT_timestamp"] = previous_distinct_in_group(
        df, "AUDITBYUSER", "ADUIT_Earliesttimestamp"
    )
    df["TImeDifference_BatchProcessing_HourFraction"] = hour_fraction_diff(
        df["EarliestBatchPROCESSINGtimestamp"], df["PreviousBatch_ProcessingTime"]
    )
    df["TImeDifference_AUDITING_HourFraction"] = hour_fraction_diff(
        df["ADUIT_Earliesttimestamp"], df["PreviousAUDIT_timestamp"]
    )
    df["CombinedStatus&InitialProb"] = combined_status(
        df["STATUSDESCRIPTION"], df["INITIALPROBLEMTYPE"]
    )
    df["Hold Duration"] = hold_duration(
        df["CombinedStatus&InitialProb"],
        df["INITIALPROBLEMDATE"],
        pd.to_datetime(df["PROCESSINGDATE"]),
        holidays,
    )

    # groups by batch -- see the BUG_COMPAT_MODE note above
    remit_key = "BATCHID"
    proc = pd.to_datetime(df["PROCESSINGDATE"])
    df["Employer_Earlist_Remit_timestamp"] = proc.groupby(
        df[remit_key], dropna=False
    ).transform("min")
    df["Employer_PreviousREMIT_Date"] = previous_distinct_in_group(
        df, "EMPLOYERNAME", "Employer_Earlist_Remit_timestamp"
    )
    _er_a, _er_b = df["Employer_Earlist_Remit_timestamp"], df["Employer_PreviousREMIT_Date"]
    if BUG_COMPAT_MODE:
        _er_a, _er_b = _er_a.fillna(DAX_BLANK_DATE), _er_b.fillna(DAX_BLANK_DATE)
    df["Employer_remit_datdiff"] = datediff_days(_er_a, _er_b).abs().astype("Int64")

    df["Weekday_ProcessingDate"] = weekday_name(proc)
    df["DaysuntilRETURN"] = days_until_return(
        df["STATUSDESCRIPTION"], df["DATRETURN_COMPLETEDATE"], proc
    )
    df["Hold Reason Categories"] = hold_reason_categories(
        df["HOLD REASON"], df["INITIALPROBLEMTYPE"], df["Hold Duration"]
    )
    return df


CREATEDBY_BUSINESS_ANALYSTS = ["MMoore", "KKing", "NNash", "JJones"]
AUDITOR_BUSINESS_ANALYSTS = ["LLane", "KKing", "JJones"]


def derive_weekly(df: pd.DataFrame, holidays: np.ndarray) -> pd.DataFrame:
    df["Current Month_EarliestBatchPROCESSINGtimestamp"] = group_min(
        df, "BATCHID", "CREATEDDATE"
    )
    df["Current Month_HR_BatchedProcessed"] = format_hour_ampm(df["CREATEDDATE"])
    df["Current Month_PreviousBatch_ProcessingTime"] = previous_distinct_in_group(
        df, "CREATEDBYUSER", "Current Month_EarliestBatchPROCESSINGtimestamp"
    )
    df["Cur_Mth_Audit_Earliesttimestamp"] = group_min(df, "BATCHID", "AUDITDATE")
    df["Current Month_PreviousAUDIT_timestamp"] = previous_distinct_in_group(
        df, "AUDITBYUSER", "Cur_Mth_Audit_Earliesttimestamp"
    )
    df["TImeDifference_BatchProcessing_HourFraction"] = hour_fraction_diff(
        df["Current Month_EarliestBatchPROCESSINGtimestamp"],
        df["Current Month_PreviousBatch_ProcessingTime"],
    )
    df["Current Month_TImeDifference_AUDITING_HourFraction"] = hour_fraction_diff(
        df["Cur_Mth_Audit_Earliesttimestamp"],
        df["Current Month_PreviousAUDIT_timestamp"],
    )
    df["CombinedStatus&InitialProb"] = combined_status(
        df["STATUSDESCRIPTION"], df["INITIALPROBLEMTYPE"]
    )
    df["Hold Duration"] = hold_duration(
        df["CombinedStatus&InitialProb"],
        df["INITIALPROBLEMDATE"],
        pd.to_datetime(df["PROCESSINGDATE"]),
        holidays,
    )
    df["DaysuntilRETURN"] = days_until_return(
        df["STATUSDESCRIPTION"],
        df["DATRETURN_COMPLETEDATE"],
        pd.to_datetime(df["PROCESSINGDATE"]),
    )
    df["Hold Reason Categories"] = hold_reason_categories(
        df["HOLD REASON"], df["INITIALPROBLEMTYPE"], df["Hold Duration"]
    )
    df["CREATEDBYUSER (groups)"] = user_group(
        df["CREATEDBYUSER"], CREATEDBY_BUSINESS_ANALYSTS, "Business Analyst"
    )
    df["Auditors"] = user_group(
        df["AUDITBYUSER"], AUDITOR_BUSINESS_ANALYSTS, "Business Analysts"
    )
    return df


DERIVERS = {"monthly": derive_monthly, "weekly": derive_weekly}


# ---------------------------------------------------------------------------
# Columns dropped by the old "Removed Columns" step. Dropping them here means
# they never enter memory at all -- this is most of the savings, more than the
# Parquet format itself.
#
# The two queries differ: the weekly one KEEPS "RETURN". Nothing else varies.
# ---------------------------------------------------------------------------
DROP_COMMON = [
    "DEPOSITID",
    "UPLOADABLEFILENAME",
    "UPLOADLOADEDDATETIME",
    "SYMETRAERCODE",
    "PRODUCT",
    "VCHRSSN",
    "PLAYERSNAME",
    "NAMEONBATCH",
    "NSCC",
    "OPP",
    "OIS",
    "SFPP",
    "SFPP_NY",
    "PREF_NONQUAL",
    "PREF_QUAL",
    "GREAT_AMERICAN",
    "NYLIFE",
    "EQUI-VEST",
    "AXA-NJ",
    "PERSHING",
    "SBL",
    "HOLD",
    "LOAN",
    "REIMBURSELOCKBOX",
    "SDBA",
    "ERCODE",
    "ACCT NO",
    "PARENTBATCH",
    "TRANCODE",
]

DATE_COLUMNS = ["PROCESSINGDATE"]
DATETIME_COLUMNS = [
    "CREATEDDATE",
    "AUDITDATE",
    "INITIALPROBLEMDATE",
    "DATRETURN_COMPLETEDATE",
]
FLOAT_COLUMNS = ["TOTAL"]
TEXT_COLUMNS = [
    "STATUSDESCRIPTION",
    "INTERNAL",
    "MODULE",
    "CREATEDBYUSER",
    "AUDITBYUSER",
    "UPLOADABLEUSED",
    "INITIALPROBLEMTYPE",
    "EMPLOYERNAME",
    "DEPOSITTRANTYPE",
    "CHILDBATCH",
    "HOLD REASON",
]

# Value cleanup, ported from the M ReplaceValue steps.
#
# M semantics matter and are easy to get wrong:
#   Replacer.ReplaceText  -> substring replace  (pandas .str.replace)
#   Replacer.ReplaceValue -> whole-value match  (pandas .replace)
EMPLOYERNAME_FIXES_MONTHLY = [
    ("SAMPLE BROTHERSVALLEY SD", "SAMPLE BROTHERS VALLEY SD"),
    ("Sample County SD", "Sample County School District"),
    ("Sample Finacial Advisors", "Sample Financial Advisors"),
    ("Sample ISD 403b Plan", "Sample ISD"),
    ("Sample ISD 457 Plan", "Sample ISD"),
    ("Sample Health 403b (SH)", "Sample Health (SH)"),
    ("Sample Health 457 (SH)", "Sample Health (SH)"),
]

MODULE_FIXES = [
    ("IndividualModule", "Individual Module"),
    ("EmployerModule", "Employer Module"),
]

# All Data: the monthly fixes plus one extra variant its old query cleaned.
# (The old query aimed the two Sample ISD fixes at Source.Name -- a no-op -- so
# they never applied; here they hit EMPLOYERNAME as intended.)
EMPLOYERNAME_FIXES_ALLDATA = EMPLOYERNAME_FIXES_MONTHLY + [
    ("Sample Health (SYHC)", "Sample Health (SH)"),
]

TRANTYPE_FIXES_WEEKLY = [
    ("TradeTicket", "Trade Ticket"),
    ("RemoteDeposit", "Remote Deposit"),
]

# The monthly query's blank -> "Internals" step on DEPOSITTRANTYPE only ever
# caught literal empty strings, because Power Query reads a blank CSV cell as
# null. Truly blank cells stayed null and fell out of any grouping. Set this
# False to reproduce that bug-for-bug while validating totals.
TREAT_NULL_TRANTYPE_AS_INTERNAL = True


@dataclass
class Profile:
    key: str
    label: str
    src: Path
    drop_columns: list
    int_columns: list
    employername_fixes: list = field(default_factory=list)
    module_fixes: list = field(default_factory=list)
    trantype_fixes: list = field(default_factory=list)
    trantype_blank_to_internals: bool = False
    # The weekly query keeps Source.Name (the filename Power Query adds during a
    # folder combine). It is not a CSV column, so we write it ourselves using the
    # original .csv name -- not the .parquet name -- so values match exactly.
    add_source_name: bool = False
    # True: read every file, derive the calculated columns across the whole
    # table, then split back per file (needed when a row depends on rows in
    # other files). False: convert file by file and skip unchanged files.
    derive: bool = True

    @property
    def out(self) -> Path:
        return self.src.parent / f"{self.src.name} (parquet)"


PROFILES = [
    Profile(
        key="monthly",
        label="MONTHLY report  (Lockbox 12month raw data)",
        src=ROOT / "Lockbox 12month raw data",
        drop_columns=DROP_COMMON + ["RETURN"],
        int_columns=["STATUSID", "BATCHID", "INVESTMENTID"],
        employername_fixes=EMPLOYERNAME_FIXES_MONTHLY,
        module_fixes=MODULE_FIXES,
        trantype_blank_to_internals=True,
        add_source_name=False,
    ),
    Profile(
        key="weekly",
        label="WEEKLY report   (Current Raw Data)",
        src=ROOT / "Current Raw Data",
        drop_columns=DROP_COMMON,          # keeps RETURN
        int_columns=["STATUSID", "BATCHID", "INVESTMENTID", "RETURN"],
        module_fixes=MODULE_FIXES,
        trantype_fixes=TRANTYPE_FIXES_WEEKLY,
        trantype_blank_to_internals=False,
        add_source_name=True,
    ),
    Profile(
        key="alldata",
        label="ALL DATA history (All Data)",
        src=ROOT / "All Data",
        drop_columns=DROP_COMMON,          # keeps RETURN
        int_columns=["STATUSID", "BATCHID", "INVESTMENTID", "RETURN"],
        employername_fixes=EMPLOYERNAME_FIXES_ALLDATA,
        module_fixes=MODULE_FIXES,
        trantype_blank_to_internals=False,
        add_source_name=True,
        # No calculated columns depend on this table, and deriving across six
        # years would mean holding ~16M rows at once. File by file instead, so
        # after the first run only a new month gets converted.
        derive=False,
    ),
]

# Folders whose data feeds the KPI Benchmarks step.
BENCH_PROFILE_KEYS = ("monthly", "weekly")


# ===========================================================================
# CONTROL SHEET -- daily UMB control sheet workbooks -> one Parquet file
# ===========================================================================
#
# The monthly report's "Control Sheet data" table comes from the daily control
# sheet workbooks (CONTROL SHEET <year>\<MM - Month>\<day>.xlsx). Each has a
# "Monthly numbers" tab: one row of formulas pointing at CONTROL SHEET.
#
# Workbooks saved from Excel store the formulas' results. Workbooks written by
# a library such as openpyxl store the formulas only -- Excel recalculates on
# open, but Power Query reads stored results and gets blanks. That's why
# September came up short. So for each workbook this step:
#   * uses Excel's stored results when present (exactly what Power Query read);
#   * otherwise evaluates that workbook's OWN formulas. On busy days the
#     trade-ticket section grows and every formula below it shifts (J2 might
#     read E48 instead of E44), so nothing here assumes the template's cell
#     addresses -- it follows whatever each file says.
# A day with a negative deposit total or an implausible count is held back and
# reported rather than passed through to the report.
# Already-read workbooks are skipped on the next run unless they changed.
# ---------------------------------------------------------------------------

from openpyxl.utils import column_index_from_string, get_column_letter

CS_ROOT = Path(r"Q:\CORE\PURCHASE TRADING\UMB Lockbox Deposits-Daily Spreadsheets")
CS_YEARS_BACK = 1                   # this year's folder plus this many before it
CS_OUT = ROOT / "Control Sheet data (parquet)" / "control_sheet.parquet"
CS_MAX_DAILY_COUNT = 2000          # any single count above this is treated as a bad read
# Bump when the reading logic changes, so cached rows from older logic are
# re-read instead of reused.
CS_ENGINE_VERSION = 2

CS_COLUMNS = ["Date", "Total  deposited", "UMB Checks", "Remote Deposit Checks",
              '"ACHs/Wires"', "TradeTickets", "ACH Debits"]   # names exactly as the report has them
_CS_EXPECTED_HEADERS = ["date", "totaldeposited", "umbchecks", "remotedepositchecks",
                        "achswires", "tradetickets", "achdebits"]


_TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<ref>(?:'(?P<sheet>[^']+)'!|(?P<bare>[A-Za-z_][\w\.]*)!)?
            \$?(?P<c1>[A-Za-z]{1,3})\$?(?P<r1>\d+)
            (?::\$?(?P<c2>[A-Za-z]{1,3})\$?(?P<r2>\d+))?)
  | (?P<num>\d+(?:\.\d+)?)
  | (?P<func>[A-Za-z]+)\(
  | (?P<op>[-+*/(),])
""", re.VERBOSE)


class FormulaError(Exception):
    pass


class _Range:
    def __init__(self, sheet, c1, r1, c2, r2):
        self.sheet, self.c1, self.r1, self.c2, self.r2 = sheet, c1, r1, c2, r2

    def coords(self):
        for r in range(self.r1, self.r2 + 1):
            for c in range(self.c1, self.c2 + 1):
                yield self.sheet, f"{get_column_letter(c)}{r}"


class WorkbookEvaluator:
    """Evaluates the small formula dialect the control sheet template uses:
    cell references (optionally sheet-qualified / absolute), ranges, numbers,
    + - * /, parentheses, SUM, COUNTA and WORKDAY(date, n) (weekends only,
    as Excel does when no holiday list is given)."""

    def __init__(self, wb):
        self.wb = wb
        self.memo = {}

    # -- cells --------------------------------------------------------------
    def raw(self, sheet, coord):
        return self.wb[sheet][coord].value

    def value(self, sheet, coord, depth=0):
        key = (sheet, coord)
        if key in self.memo:
            return self.memo[key]
        if depth > 200:
            raise FormulaError("reference chain too deep")
        v = self.raw(sheet, coord)
        if isinstance(v, str) and v.startswith("="):
            v = self.formula(v[1:], sheet, depth + 1)
        self.memo[key] = v
        return v

    # -- parsing ------------------------------------------------------------
    def formula(self, text, sheet, depth=0):
        # Evaluating a referenced cell parses *its* formula mid-way through
        # this one, so the parser state is saved and restored around it.
        saved = (getattr(self, "_toks", None), getattr(self, "_i", 0),
                 getattr(self, "_sheet", None), getattr(self, "_depth", 0))
        try:
            self._toks = self._tokenize(text, sheet)
            self._i, self._sheet, self._depth = 0, sheet, depth
            out = self._expr()
            if self._i != len(self._toks):
                raise FormulaError(f"unparsed: {text}")
            toks, i = self._toks, self._i
            result = self._scalar(out)
            return result
        finally:
            self._toks, self._i, self._sheet, self._depth = saved

    def _tokenize(self, text, sheet):
        toks, pos = [], 0
        while pos < len(text):
            m = _TOKEN.match(text, pos)
            if not m:
                raise FormulaError(f"can't read formula at: {text[pos:]}")
            pos = m.end()
            if m.group("ws"):
                continue
            if m.group("ref"):
                sh = m.group("sheet") or m.group("bare") or sheet
                c1, r1 = column_index_from_string(m.group("c1").upper()), int(m.group("r1"))
                if m.group("c2"):
                    c2, r2 = column_index_from_string(m.group("c2").upper()), int(m.group("r2"))
                else:
                    c2, r2 = c1, r1
                toks.append(("ref", _Range(sh, c1, r1, c2, r2)))
            elif m.group("num"):
                toks.append(("num", float(m.group("num"))))
            elif m.group("func"):
                toks.append(("func", m.group("func").upper()))
            else:
                toks.append(("op", m.group("op")))
        return toks

    def _peek(self):
        return self._toks[self._i] if self._i < len(self._toks) else (None, None)

    def _take(self, kind=None, val=None):
        t = self._peek()
        if (kind and t[0] != kind) or (val and t[1] != val):
            raise FormulaError(f"expected {val or kind}, got {t[1]!r}")
        self._i += 1
        return t

    def _expr(self):
        left = self._term()
        while self._peek() in (("op", "+"), ("op", "-")):
            op = self._take()[1]
            a, b = self._num(left), self._num(self._term())
            left = a + b if op == "+" else a - b
        return left

    def _term(self):
        left = self._factor()
        while self._peek() in (("op", "*"), ("op", "/")):
            op = self._take()[1]
            a, b = self._num(left), self._num(self._factor())
            left = a * b if op == "*" else a / b
        return left

    def _factor(self):
        kind, val = self._peek()
        if (kind, val) == ("op", "-"):
            self._take(); return -self._num(self._factor())
        if (kind, val) == ("op", "+"):
            self._take(); return self._num(self._factor())
        if kind == "num":
            self._take(); return val
        if kind == "ref":
            self._take(); return val
        if (kind, val) == ("op", "("):
            self._take(); v = self._expr(); self._take("op", ")"); return v
        if kind == "func":
            self._take()
            args = []
            if self._peek() != ("op", ")"):
                args.append(self._expr())
                while self._peek() == ("op", ","):
                    self._take(); args.append(self._expr())
            self._take("op", ")")
            return self._call(val, args)
        raise FormulaError(f"unexpected {val!r}")

    # -- values -------------------------------------------------------------
    def _scalar(self, v):
        if isinstance(v, _Range):
            if (v.c1, v.r1) != (v.c2, v.r2):
                raise FormulaError("range used as a single value")
            return self.value(v.sheet, f"{get_column_letter(v.c1)}{v.r1}", self._depth)
        return v

    def _num(self, v):
        v = self._scalar(v)
        if v is None:
            return 0.0
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            raise FormulaError(f"not a number: {v!r}")
        return float(v)

    def _call(self, name, args):
        if name == "SUM":
            total = 0.0
            for a in args:
                if isinstance(a, _Range):
                    for sh, co in a.coords():
                        v = self.value(sh, co, self._depth)
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            total += v
                else:
                    total += self._num(a)
            return total
        if name == "COUNTA":
            n = 0
            for a in args:
                if isinstance(a, _Range):
                    n += sum(1 for sh, co in a.coords() if self.raw(sh, co) not in (None, ""))
                else:
                    n += 1
            return float(n)
        if name == "WORKDAY":
            d, k = self._scalar(args[0]), int(self._num(args[1]))
            if not isinstance(d, datetime.datetime):
                raise FormulaError("WORKDAY needs a date")
            step = 1 if k >= 0 else -1
            for _ in range(abs(k)):
                d += datetime.timedelta(days=step)
                while d.weekday() >= 5:
                    d += datetime.timedelta(days=step)
            return d
        raise FormulaError(f"unsupported function {name}")


def _cs_headers_ok(ws) -> bool:
    headers = [re.sub(r"[^a-z0-9]", "", str(c.value or "").lower()) for c in ws[1]][:7]
    return headers == _CS_EXPECTED_HEADERS


def cs_stored(wb):
    """Excel's stored results from 'Monthly numbers', or None if there aren't any."""
    if "Monthly numbers" not in wb.sheetnames or not _cs_headers_ok(wb["Monthly numbers"]):
        return None
    vals = [c.value for c in wb["Monthly numbers"][2]][:7]
    if not isinstance(vals[0], datetime.datetime):
        return None
    vals[0] = vals[0].date()
    return vals


def cs_evaluate(wb):
    """Evaluate the workbook's own 'Monthly numbers' formulas. Returns (values, error)."""
    if "Monthly numbers" not in wb.sheetnames or not _cs_headers_ok(wb["Monthly numbers"]):
        return None, "no recognisable 'Monthly numbers' tab"
    try:
        ev = WorkbookEvaluator(wb)
        vals = [ev.value("Monthly numbers", f"{c}2") for c in "ABCDEFG"]
    except FormulaError as exc:
        return None, f"couldn't evaluate its formulas: {exc}"
    if not isinstance(vals[0], datetime.datetime):
        return None, "its date didn't evaluate to a date"
    vals[0] = vals[0].date()
    vals[1:] = [None if v is None else float(v) for v in vals[1:]]
    vals[1] = None if vals[1] is None else round(vals[1], 2)
    return vals, None


def _cs_differs(a, b) -> bool:
    for x, y in zip(a, b):
        if isinstance(x, datetime.date) or isinstance(y, datetime.date):
            if x != y:
                return True
        elif abs(float(x or 0) - float(y or 0)) > 0.005:
            return True
    return False


def _cs_read(path: Path):
    """Returns (row values, method, note)."""
    import openpyxl
    stored = cs_stored(openpyxl.load_workbook(path, data_only=True))
    computed, err = cs_evaluate(openpyxl.load_workbook(path, data_only=False))
    if stored is not None:
        note = None
        if computed is not None and _cs_differs(computed, stored):
            note = "Excel's stored values differ from its formulas -- Excel's used"
        return stored, "excel", note
    if computed is not None:
        return computed, "computed", None
    return None, "skipped", err


def run_control_sheet(args) -> int:
    print("=" * 74)
    print("CONTROL SHEET  (daily UMB workbooks -> Control Sheet data)")
    print("=" * 74)
    print(f"  source : {CS_ROOT}")
    print(f"  output : {CS_OUT}\n")
    if not CS_ROOT.is_dir():
        print("  ERROR: control sheet folder not found.\n", file=sys.stderr)
        return 1

    this_year = datetime.date.today().year
    files = []
    for year in range(this_year - CS_YEARS_BACK, this_year + 1):
        folder = CS_ROOT / f"CONTROL SHEET {year}"
        if folder.is_dir():
            files += [p for p in folder.rglob("*.xlsx") if not p.name.startswith("~$")]
    if not files:
        print("  ERROR: no control sheet workbooks found.\n", file=sys.stderr)
        return 1

    cache = {}
    if CS_OUT.exists() and not args.force:
        try:
            old = pd.read_parquet(CS_OUT)
            for rec in old.to_dict("records"):
                if rec.get("__engine") == CS_ENGINE_VERSION:
                    cache[(rec["__path"], rec["__mtime"])] = rec
        except Exception as exc:
            print(f"  note: couldn't read previous output ({exc}) -- rebuilding")

    records, counts, notes = [], {"reused": 0, "excel": 0, "computed": 0, "skipped": 0, "failed": 0}, []
    start = time.perf_counter()
    for p in sorted(files):
        key = (str(p), p.stat().st_mtime)
        if key in cache:
            records.append(cache[key])
            counts["reused"] += 1
            continue
        try:
            vals, method, note = _cs_read(p)
        except Exception as exc:
            counts["failed"] += 1
            notes.append(f"{p.parent.name}\\{p.name}: FAILED {type(exc).__name__}: {exc}")
            continue
        counts[method] += 1
        if note:
            notes.append(f"{p.parent.name}\\{p.name}: {note}")
        if vals is None:
            continue
        rec = dict(zip(CS_COLUMNS, vals))
        rec.update({"Source.Name": p.name, "__path": key[0], "__mtime": key[1],
                    "__method": method, "__engine": CS_ENGINE_VERSION})
        records.append(rec)

    print(f"  {len(files)} workbooks: {counts['reused']} unchanged since last run, "
          f"{counts['excel']} read from Excel's stored values, "
          f"{counts['computed']} computed (no stored values), "
          f"{counts['skipped']} skipped, {counts['failed']} failed  "
          f"[{time.perf_counter() - start:.1f}s]")
    for n in notes[:15]:
        print(f"    {n}")
    if len(notes) > 15:
        print(f"    ... and {len(notes) - 15} more")
    if not records:
        print("  ERROR: nothing usable was read.\n", file=sys.stderr)
        return 1

    df = pd.DataFrame(records)
    df["Date"] = pd.to_datetime(df["Date"]).dt.date

    # Guard: hold back days that can't be right rather than chart them.
    counts_cols = ["UMB Checks", "Remote Deposit Checks", '"ACHs/Wires"', "TradeTickets", "ACH Debits"]
    num = df[counts_cols].apply(pd.to_numeric, errors="coerce")
    total = pd.to_numeric(df["Total  deposited"], errors="coerce")
    bad = (total < 0) | (num > CS_MAX_DAILY_COUNT).any(axis=1) | (num < 0).any(axis=1)
    if bad.any():
        print(f"  WARNING: {int(bad.sum())} day(s) held back -- values can't be right:")
        for _, r in df[bad].iterrows():
            rel = Path(r["__path"]).relative_to(CS_ROOT) if str(r["__path"]).startswith(str(CS_ROOT)) else r["__path"]
            print(f"    {rel}  ({r['Date']}): total {r['Total  deposited']}, "
                  + ", ".join(f"{c} {r[c]}" for c in counts_cols))
        df = df[~bad]
    # One row per day: if a date appears in more than one workbook, the
    # most recently saved one wins.
    before = len(df)
    df = df.sort_values("__mtime").drop_duplicates("Date", keep="last").sort_values("Date")
    if len(df) < before:
        print(f"  {before - len(df)} duplicate date(s) -- kept the most recently saved workbook")
    for c in ("UMB Checks", "Remote Deposit Checks", '"ACHs/Wires"', "TradeTickets", "ACH Debits"):
        df[c] = pd.to_numeric(df[c], errors="coerce").round().astype("Int64")
    df["Total  deposited"] = pd.to_numeric(df["Total  deposited"], errors="coerce").astype("Float64")
    df = df[["Source.Name"] + CS_COLUMNS + ["__path", "__mtime", "__method", "__engine"]]

    CS_OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(CS_OUT, engine="pyarrow", index=False)

    months = pd.Series(pd.to_datetime(df["Date"])).dt.to_period("M").value_counts().sort_index()
    print(f"  wrote {len(df):,} days, {df['Date'].min()} .. {df['Date'].max()}")
    print("  days per month (recent): " + ",  ".join(f"{m}: {n}" for m, n in months.tail(4).items()) + "\n")
    return 1 if counts["failed"] else 0



def read_csv_safe(csv_path: Path, report: list, **kwargs) -> pd.DataFrame:
    """Read a CSV, falling back to cp1252 if it isn't valid UTF-8.

    These exports often come out as Windows-1252, in which case one accented
    employer name would kill a utf-8 read.
    """
    try:
        return pd.read_csv(csv_path, encoding="utf-8-sig", **kwargs)
    except UnicodeDecodeError:
        report.append("    note: not valid UTF-8, re-read as cp1252")
        return pd.read_csv(csv_path, encoding="cp1252", **kwargs)


def normalize(df: pd.DataFrame, prof: Profile) -> pd.DataFrame:
    if prof.employername_fixes and "EMPLOYERNAME" in df.columns:
        for old, new in prof.employername_fixes:
            df["EMPLOYERNAME"] = df["EMPLOYERNAME"].str.replace(old, new, regex=False)

    if prof.module_fixes and "MODULE" in df.columns:
        for old, new in prof.module_fixes:
            df["MODULE"] = df["MODULE"].str.replace(old, new, regex=False)

    if prof.trantype_fixes and "DEPOSITTRANTYPE" in df.columns:
        for old, new in prof.trantype_fixes:
            df["DEPOSITTRANTYPE"] = df["DEPOSITTRANTYPE"].str.replace(old, new, regex=False)

    if prof.trantype_blank_to_internals and "DEPOSITTRANTYPE" in df.columns:
        df["DEPOSITTRANTYPE"] = df["DEPOSITTRANTYPE"].replace("", "Internals")
        if TREAT_NULL_TRANTYPE_AS_INTERNAL:
            df["DEPOSITTRANTYPE"] = df["DEPOSITTRANTYPE"].fillna("Internals")

    return df


def coerce_types(df: pd.DataFrame, prof: Profile, report: list) -> pd.DataFrame:
    for col in DATE_COLUMNS:
        if col in df.columns:
            before = df[col].notna().sum()
            parsed = pd.to_datetime(df[col], errors="coerce")
            lost = before - parsed.notna().sum()
            if lost:
                report.append(f"    {col}: {lost} value(s) failed date parse -> null")
            # .dt.date yields object dtype, which pyarrow writes as date32. An
            # all-null column gives it nothing to infer from, so leave those as
            # timestamps and let Power Query cast.
            df[col] = parsed.dt.date if parsed.notna().any() else parsed

    for col in DATETIME_COLUMNS:
        if col in df.columns:
            before = df[col].notna().sum()
            parsed = pd.to_datetime(df[col], errors="coerce")
            lost = before - parsed.notna().sum()
            if lost:
                report.append(f"    {col}: {lost} value(s) failed datetime parse -> null")
            df[col] = parsed

    for col in prof.int_columns:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").astype("Int64")

    for col in FLOAT_COLUMNS:
        if col in df.columns:
            cleaned = (
                df[col]
                .astype("string")
                .str.replace(r"[$,]", "", regex=True)
                .str.replace(r"^\((.*)\)$", r"-\1", regex=True)  # (123.45) -> -123.45
                .str.strip()
            )
            df[col] = pd.to_numeric(cleaned, errors="coerce").astype("Float64")

    for col in TEXT_COLUMNS:
        if col in df.columns:
            df[col] = df[col].astype("string")

    return df


def clean_one(csv_path: Path, prof: Profile):
    """Read and clean one CSV. Does NOT write -- the derived columns need the
    whole table, so writing happens after every file is concatenated."""
    report: list = []

    df = read_csv_safe(
        csv_path,
        report,
        dtype="string",          # read as text, cast deliberately below
        keep_default_na=True,
        na_values=[""],
        low_memory=False,
    )

    df.columns = [c.strip() for c in df.columns]

    present = [c for c in prof.drop_columns if c in df.columns]
    missing = [c for c in prof.drop_columns if c not in df.columns]
    if missing:
        shown = ", ".join(missing[:6])
        extra = f" (+{len(missing) - 6} more)" if len(missing) > 6 else ""
        report.append(f"    note: {len(missing)} drop-list column(s) absent: {shown}{extra}")
    df = df.drop(columns=present)

    df = normalize(df, prof)
    df = coerce_types(df, prof, report)

    if prof.add_source_name:
        df.insert(0, "Source.Name", pd.Series([csv_path.name] * len(df), dtype="string"))

    return df, report


def audit(prof: Profile, csvs: list) -> bool:
    """Compare column sets across CSVs without loading the data.

    Worth running before switching a report over. The old queries expanded
    using the *sample file's* column list, so a file with extra columns
    silently lost them and a file missing columns silently got nulls -- with
    no warning either way.

    Returns True if drift was found.
    """
    schemas = {}
    rows = {}

    for path in csvs:
        report: list = []
        head = read_csv_safe(path, report, nrows=0, dtype="string")
        schemas[path] = [c.strip() for c in head.columns]
        with path.open("rb") as fh:
            rows[path] = max(sum(1 for _ in fh) - 1, 0)
        for line in report:
            print(f"  {path.name}{line}")

    reference = schemas[csvs[0]]
    ref_set = set(reference)
    print(f"  reference: {csvs[0].name}  ({len(reference)} columns)\n")

    drift = False
    for path in csvs:
        cols = schemas[path]
        extra = [c for c in cols if c not in ref_set]
        absent = [c for c in reference if c not in set(cols)]
        flags = []
        if extra:
            flags.append(f"EXTRA: {', '.join(extra)}")
        if absent:
            flags.append(f"MISSING: {', '.join(absent)}")
        if not flags and cols != reference:
            flags.append("column order differs (harmless)")
        if extra or absent:
            drift = True
        print(
            f"  {path.name[:52]:52} {len(cols):3} cols  ~{rows[path]:>9,} rows  "
            f"{'; '.join(flags) if flags else 'matches'}"
        )
    return drift


def _convert_per_file(prof: Profile, csvs: list, args, compression) -> int:
    """File-by-file conversion for profiles with no derived columns."""
    built = skipped = failed = 0
    csv_bytes = pq_bytes = 0
    for csv_path in csvs:
        parquet_path = prof.out / (csv_path.stem + ".parquet")
        csv_bytes += csv_path.stat().st_size
        if (not args.force and parquet_path.exists()
                and parquet_path.stat().st_mtime >= csv_path.stat().st_mtime):
            skipped += 1
            pq_bytes += parquet_path.stat().st_size
            continue
        start = time.perf_counter()
        try:
            df, report = clean_one(csv_path, prof)
            df.to_parquet(parquet_path, engine="pyarrow", compression=compression, index=False)
        except Exception as exc:  # one bad file shouldn't stop the run
            print(f"  FAILED {csv_path.name}: {type(exc).__name__}: {exc}", file=sys.stderr)
            failed += 1
            continue
        pq_bytes += parquet_path.stat().st_size
        print(f"  build  {csv_path.name}  {len(df):,} rows, "
              f"{parquet_path.stat().st_size / 1024 / 1024:.1f} MB, "
              f"{time.perf_counter() - start:.1f}s")
        for line in report:
            print(line)
        built += 1
    if skipped:
        print(f"  skip   {skipped} file(s) already up to date")

    if not args.keep_orphans:
        expected = {p.stem for p in csvs}
        for stale in sorted(prof.out.glob("*.parquet")):
            if stale.stem not in expected:
                print(f"  remove {stale.name}  (source CSV no longer present)")
                stale.unlink()

    print(f"\n  {built} built, {skipped} up to date, {failed} failed")
    print(f"  {csv_bytes / 1024 / 1024:.0f} MB CSV -> {pq_bytes / 1024 / 1024:.0f} MB Parquet\n")
    return 1 if failed else 0


def run_profile(prof: Profile, args) -> int:
    print("=" * 74)
    print(prof.label)
    print("=" * 74)
    print(f"  source : {prof.src}")
    print(f"  output : {prof.out}\n")

    if not prof.src.is_dir():
        print("  ERROR: source folder not found.\n", file=sys.stderr)
        return 1

    csvs = sorted(p for p in prof.src.glob("*.csv") if not p.name.startswith("~"))
    if not csvs:
        print("  ERROR: no CSV files found.\n", file=sys.stderr)
        return 1

    if args.audit:
        drift = audit(prof, csvs)
        print()
        return 1 if drift else 0

    prof.out.mkdir(parents=True, exist_ok=True)
    compression = None if args.compression == "none" else args.compression

    if not prof.derive:
        return _convert_per_file(prof, csvs, args, compression)

    # The derived columns need the whole table: PreviousBatch_ProcessingTime
    # looks for a user's previous batch, which usually sits in the prior
    # month's file. So everything is read, concatenated, derived, then split
    # back out per source file. That means the skip check is all-or-nothing --
    # a new month changes neighbouring rows, so a partial rebuild would be wrong.
    newest_csv = max(p.stat().st_mtime for p in csvs)
    parquets = [prof.out / (p.stem + ".parquet") for p in csvs]
    version_file = prof.out / "_derive_version.txt"
    built_version = version_file.read_text().strip() if version_file.exists() else None
    if built_version != str(DERIVE_VERSION):
        print(f"  derived-column logic changed (built with {built_version or 'none'}, "
              f"now {DERIVE_VERSION}) -- rebuilding")
    elif not args.force and all(q.exists() for q in parquets) and all(
        q.stat().st_mtime >= newest_csv for q in parquets
    ):
        print(f"  all {len(csvs)} file(s) up to date -- nothing to do.\n")
        return 0

    try:
        holidays = load_holidays(args.holidays)
    except Exception as exc:
        print(f"  ERROR: could not read holidays from {args.holidays}: {exc}\n", file=sys.stderr)
        return 1

    frames, failed = [], 0
    csv_bytes = 0
    for csv_path in csvs:
        print(f"  read   {csv_path.name}")
        try:
            df, report = clean_one(csv_path, prof)
        except Exception as exc:
            print(f"    FAILED: {type(exc).__name__}: {exc}", file=sys.stderr)
            failed += 1
            continue
        for line in report:
            print(line)
        df["__source"] = csv_path.stem
        frames.append(df)
        csv_bytes += csv_path.stat().st_size

    if not frames:
        print("  ERROR: nothing read successfully.\n", file=sys.stderr)
        return 1

    combined = pd.concat(frames, ignore_index=True)
    del frames
    print(f"\n  deriving {len(combined):,} rows across {len(csvs)} file(s)...")
    start = time.perf_counter()
    combined = DERIVERS[prof.key](combined, holidays)
    print(f"  derived in {time.perf_counter() - start:.1f}s\n")

    pq_bytes = 0
    for csv_path, parquet_path in zip(csvs, parquets):
        part = combined[combined["__source"] == csv_path.stem].drop(columns="__source")
        part.to_parquet(parquet_path, engine="pyarrow", compression=compression, index=False)
        out_mb = parquet_path.stat().st_size / 1024 / 1024
        pq_bytes += parquet_path.stat().st_size
        print(f"  write  {parquet_path.name}  {len(part):,} rows, {out_mb:.1f} MB")
    if not failed:
        version_file.write_text(str(DERIVE_VERSION))

    if not args.keep_orphans:
        expected = {p.stem for p in csvs}
        for stale in sorted(prof.out.glob("*.parquet")):
            if stale.stem not in expected:
                print(f"  remove {stale.name}  (source CSV no longer present)")
                stale.unlink()

    print(f"\n  {len(csvs) - failed} file(s) written, {failed} failed, "
          f"{len(combined.columns)} columns")
    print(f"  {csv_bytes / 1024 / 1024:.0f} MB CSV -> {pq_bytes / 1024 / 1024:.0f} MB Parquet\n")
    return 1 if failed else 0


# ===========================================================================
# KPI BENCHMARKS -- recompute the current month's row
# ===========================================================================
#
# Each row of KPI Benchmarks.xlsx is a trailing window: the "Year Range" label
# on the August 2026 row reads "08/25 to 08/26", i.e. 13 monthly files, M-12
# through M. The current month's row is the 12-month folder (minus its oldest
# month) plus the month-to-date file in Current Raw Data.
#
# Rates are batches per hour. The weekly report's DAX is:
#
#     Rate = DISTINCTCOUNT(BATCHID)
#            / [sum of hour-fraction gaps, unique batches, under 15 minutes]
#
# where the "gap" is the time since the same user's previous batch
# (TImeDifference_BatchProcessing_HourFraction). Gaps of 15 minutes or more are
# treated as breaks and contribute no time. Auditing works the same way using
# the audit timestamps and AUDITBYUSER.
#
# The DAX behind the 12-month benchmark measures isn't fully known, so the
# script CALIBRATES: it tries every plausible variant of the numerator and
# denominator against historical benchmark rows it has data for, picks the
# variant that reproduces them best, and only writes the new row if that fit is
# within BENCH_FIT_TOLERANCE. Paste the calibration output back if it misses.
# ---------------------------------------------------------------------------

KPI_XLSX = ROOT / "KPI Benchmarks.xlsx"
KPI_BACKUP_DIR = ROOT / "KPI Benchmarks Backups"
BENCH_WINDOW_MONTHS = 13            # "08/25 to 08/26" inclusive
HF_THRESHOLD_HOURS = 0.25           # "Under15Min"
BENCH_FIT_TOLERANCE = 0.03          # mean abs % error allowed before writing
BENCH_MIN_CALIB_COVERAGE = 0.75     # calibrate on rows with >= 75% of window data

# How the new row's values are produced:
#   "chain"  new = previous month's SHEET value x (computed this month / computed
#            last month). The level carries on from the sheet; the month-over-
#            month movement comes from the real data. Used because 10 columns
#            (the uploadable splits and the auditing module splits) follow DAX
#            we haven't reproduced, and some past rows are estimates -- so exact
#            matching isn't the goal, continuity with the trend is. For columns
#            whose definition does match, the link factor is ~1 anyway.
#   "raw"    new = computed value as-is. Only written if calibration matches
#            within BENCH_FIT_TOLERANCE (or with --force-benchmarks).
BENCH_STRATEGY = "chain"

# Definitions confirmed from the report's DAX. When set, these are used instead
# of whichever calibration variant sits closest to the sheet's levels -- under
# chain mode only month-over-month movement matters, so the real definition
# beats the best level fit. Set a family to None to fall back to best fit.
#   P: Cur_Mth_Distinct_BATCHIDS / CurMth_Total_HourFraction_UniqueBatches_Under15Min,
#      segment measures filtered CREATEDBYUSER <> "SYSTEM" && <> ""
#   A: Cur_Mth_Audits_DistinctBatch_Count_NonSystem_V2 /
#      CurMth_Total_HourFraction_UniqueAUDITS_Under15Min (AUDITBYUSER <> "SYSTEM")
BENCH_PINNED_VARIANTS = {
    "P": dict(num="all", den="batch", creator="segments", excl=False),
    "A": dict(num="auditbase", den="batch", creator="none", excl=True),
}

# Columns written as the raw computed value instead of chain-linked. The old
# 12-month measures for these used || where they meant && -- e.g.
#   [DEPOSITTRANTYPE] = "ACH" || [UPLOADABLEUSED] = "Yes"
# -- so the sheet's history for them measures "ACH OR uploadable", while the
# weekly bars drawn against them use AND. Linking would carry that wrong level
# forward; writing the raw value switches them to the correct definition. The
# first run after the switch flags them as a >10% move -- that's the fix.
BENCH_RAW_COLUMNS = {
    "ACH w/ Uploadable - Processing",
    "ACH - Processing W/O Uploadable",
    "Wire - Processing W/ Uploadable",
    "Wire - Processing W/O Uploadable",
}

# Flag any column whose new value moves more than this vs the previous month.
BENCH_JUMP_WARN = 0.10

# (sheet header, family, deposit type, module, uploadable)
#   family P = processing, A = auditing
#   uploadable: None = any, "with" / "without"
BENCH_COLUMNS = [
    ("Processing Rate",                               "P", None,     None,         None),
    ("ACH - Processing",                              "P", "ach",    None,         None),
    ("ACH w/ Uploadable - Processing",                "P", "ach",    None,         "with"),
    ("ACH - Processing W/O Uploadable",               "P", "ach",    None,         "without"),
    ("Checks Only - Processing",                      "P", "check",  None,         None),
    ("Remote Deposit - Processing",                   "P", "remote", None,         None),
    ("Trade Ticket - Processing",                     "P", "trade",  None,         None),
    ("Wire - Processing",                             "P", "wire",   None,         None),
    ("Wire - Processing W/ Uploadable",               "P", "wire",   None,         "with"),
    ("Wire - Processing W/O Uploadable",              "P", "wire",   None,         "without"),
    ("Check Only: Employer Module - Processing",      "P", "check",  "employer",   None),
    ("Check Only: Individual Module - Processing",    "P", "check",  "individual", None),
    ("Auditing Rate",                                 "A", None,     None,         None),
    ("ACH - Auditing",                                "A", "ach",    None,         None),
    ("ACH w/ Uploadable - Auditing",                  "A", "ach",    None,         "with"),
    ("ACH - Auditing W/O Uploadable",                 "A", "ach",    None,         "without"),
    ("Checks Only - Auditing",                        "A", "check",  None,         None),
    ("Remote Deposit - Auditing",                     "A", "remote", None,         None),
    ("Trade Ticket - Auditing",                       "A", "trade",  None,         None),
    ("Wire - Auditing",                               "A", "wire",   None,         None),
    ("Wire w/ Uploadable - Auditing",                 "A", "wire",   None,         "with"),
    ("Wire - Auditing W/O Uploadable",                "A", "wire",   None,         "without"),
    ("Checks Only: Employer Module - Auditing",       "A", "check",  "employer",   None),
    ("Checks Only: Individual Module - Auditing",     "A", "check",  "individual", None),
]
# "Wire w/ Uploadable Usage" is deliberately absent: it has never held a value
# and nothing defines it. The script leaves that column alone.

FAMILY_LABEL = {"P": "PROCESSING", "A": "AUDITING"}

_TRANTYPE_KEYS = {
    "ach": "ach", "checkonly": "check", "checksonly": "check", "check": "check",
    "checks": "check", "remotedeposit": "remote", "tradeticket": "trade", "wire": "wire",
}
_MODULE_KEYS = {"individualmodule": "individual", "employermodule": "employer"}
_TRUTHY = {"y", "yes", "true", "t", "1"}
_FALSY = {"n", "no", "false", "f", "0"}
_AUDIT_EXCLUDED_STATUSES = {
    "Delete", "Return Complete - Money Returned", "Return Complete - Unprocessed",
}

# Variants tried during calibration. Names appear in the output.
#   num  all        DISTINCTCOUNT(BATCHID) over every batch   (confirmed by DAX)
#        under      only batches whose own gap is under 15 min
#        auditbase  the weekly report's Audits_DistinctBatch_Count measure
#   den  batch      SUMX(FILTER(VALUES(BATCHID), MAX(gap) < 0.25), MAX(gap))
#                   -- each batch's largest gap, kept if under 15 min
#                   (confirmed by CurMth_Total_HourFraction_UniqueBatches_Under15Min)
#        row        sum of gaps over every row (investment)
#        distinct   sum of DISTINCT gap values -- a common DAX slip
#   creator  none      no CREATEDBYUSER filter
#            segments  CREATEDBYUSER <> "SYSTEM" && <> "" on every column EXCEPT the
#                      overall rate -- how the weekly segment measures are written,
#                      while the weekly overall rate measure has no such filter
#            all       that filter on every column including the overall rate
#   excl (auditing only) AUDITBYUSER = "SYSTEM" excluded from both sides
BENCH_VARIANTS = {
    "P": [
        dict(num=n, den=d, creator=cr, excl=False)
        for n in ("all", "under") for d in ("batch", "row", "distinct")
        for cr in ("none", "segments", "all")
    ],
    "A": [
        dict(num=n, den=d, creator=cr, excl=e)
        for n in ("all", "under", "auditbase") for d in ("batch", "row", "distinct")
        for cr in ("none", "segments", "all") for e in (False, True)
    ],
}
# W/O Uploadable: "explicit" = UPLOADABLEUSED is a no-value; "not_with" = anything
# that isn't a yes (so blanks count as without). Chosen separately.
WO_MODES = ("explicit", "not_with")


def _month_of(stem: str):
    m = re.match(r"^(\d{4})-(\d{1,2})(?!\d)", stem)
    return pd.Period(year=int(m.group(1)), month=int(m.group(2)), freq="M") if m else None


def _norm_key(s: pd.Series) -> pd.Series:
    return s.astype("string").str.lower().str.replace(r"[^a-z0-9]", "", regex=True)


def _window(end: pd.Period) -> list:
    return [end - i for i in range(BENCH_WINDOW_MONTHS - 1, -1, -1)]


def _range_label(end: pd.Period) -> str:
    start = end - (BENCH_WINDOW_MONTHS - 1)
    return f"{start.month:02d}/{start.year % 100:02d} to {end.month:02d}/{end.year % 100:02d}"


_BENCH_READ_COLS = [
    "BATCHID", "CREATEDDATE", "CREATEDBYUSER", "AUDITDATE", "AUDITBYUSER",
    "DEPOSITTRANTYPE", "MODULE", "UPLOADABLEUSED", "STATUSDESCRIPTION", "INITIALPROBLEMTYPE",
]


def load_bench_data():
    """Read the Parquet outputs of both folders. Returns (frame, current_month)."""
    frames, current_months = [], set()
    for prof in PROFILES:
        if prof.key not in BENCH_PROFILE_KEYS:
            continue
        for pq in sorted(prof.out.glob("*.parquet")):
            month = _month_of(pq.stem)
            if month is None:
                print(f"  skip   {pq.name}  (can't read a YYYY-M month from the name)")
                continue
            df = pd.read_parquet(pq, columns=_BENCH_READ_COLS)
            df["__month"] = month
            df["__src"] = prof.key
            frames.append(df)
            if prof.key == "weekly":
                current_months.add(month)
    if not frames:
        return None, None
    data = pd.concat(frames, ignore_index=True)
    del frames
    # Repeated text columns as categories: roughly a third of the memory.
    for c in ("CREATEDBYUSER", "AUDITBYUSER", "DEPOSITTRANTYPE", "MODULE",
              "UPLOADABLEUSED", "STATUSDESCRIPTION", "INITIALPROBLEMTYPE", "__src"):
        data[c] = data[c].astype("category")
    # A month present in both folders: the Current Raw Data copy is newer.
    if current_months:
        stale = (data["__src"] == "monthly") & data["__month"].isin(current_months)
        if stale.any():
            print(f"  note: {int(stale.sum()):,} monthly rows superseded by Current Raw Data")
        data = data[~stale].reset_index(drop=True)
    current = max(current_months) if current_months else None
    return data, current


def build_bench_frames(data: pd.DataFrame) -> dict:
    """Collapse rows to one per distinct (batch, attributes, gap) with a row count.

    Gaps are computed over everything loaded, exactly like the model's
    calculated columns are computed over the whole table.
    """
    t_raw = _norm_key(data["DEPOSITTRANTYPE"])
    data["tkey"] = t_raw.map(lambda k: _TRANTYPE_KEYS.get(k, k) if isinstance(k, str) else "<blank>")
    m_raw = _norm_key(data["MODULE"])
    data["mkey"] = m_raw.map(lambda k: _MODULE_KEYS.get(k, k) if isinstance(k, str) else "<blank>")
    u = data["UPLOADABLEUSED"].astype("string").str.strip().str.lower()
    data["upl"] = np.where(u.isin(_TRUTHY), "Y", np.where(u.isin(_FALSY), "N", "?"))

    status = data["STATUSDESCRIPTION"].astype("string").str.strip()
    ipt = data["INITIALPROBLEMTYPE"].astype("string").str.strip()
    auditor = data["AUDITBYUSER"].astype("string")

    data["deleted"] = (status == "Delete").fillna(False).astype(bool)
    # DAX: CREATEDBYUSER <> "SYSTEM" && CREATEDBYUSER <> ""  (blank = "" in DAX)
    creator = data["CREATEDBYUSER"].astype("string").str.strip()
    data["creator_bad"] = (creator.isna() | (creator == "") | (creator == "SYSTEM")).fillna(True).astype(bool)
    data["audited"] = data["AUDITDATE"].notna()
    data["system"] = (auditor == "SYSTEM").fillna(False).astype(bool)
    # Weekly Audits_DistinctBatch_Count: DAX "<>" treats blank as passing.
    data["abase_main"] = (
        data["audited"]
        & ~status.isin(_AUDIT_EXCLUDED_STATUSES).fillna(False)
        & ~(ipt == "Hold").fillna(False)
        & ~data["system"]
    )
    data["abase_hold"] = ((status == "Processed") & (ipt == "Hold")).fillna(False).astype(bool)

    tmp = pd.DataFrame({"BATCHID": data["BATCHID"], "CREATEDBYUSER": data["CREATEDBYUSER"],
                        "AUDITBYUSER": data["AUDITBYUSER"]})
    tmp["p_first"] = group_min(data, "BATCHID", "CREATEDDATE")
    tmp["p_prev"] = previous_distinct_in_group(tmp, "CREATEDBYUSER", "p_first")
    data["hf_P"] = hour_fraction_diff(tmp["p_first"], tmp["p_prev"])
    tmp["a_first"] = group_min(data, "BATCHID", "AUDITDATE")
    tmp["a_prev"] = previous_distinct_in_group(tmp, "AUDITBYUSER", "a_first")
    data["hf_A"] = hour_fraction_diff(tmp["a_first"], tmp["a_prev"])
    del tmp

    common = ["BATCHID", "__month", "tkey", "mkey", "upl", "deleted", "creator_bad"]
    frames = {}
    for fam, extra in (("P", []), ("A", ["audited", "system", "abase_main", "abase_hold"])):
        cols = common + extra + [f"hf_{fam}"]
        g = (data[cols].groupby(cols, dropna=False, observed=True).size()
             .reset_index(name="n").rename(columns={f"hf_{fam}": "hf"}))
        g["hf"] = g["hf"].astype("float64")
        frames[fam] = g
    return frames


def _seg_mask(g: pd.DataFrame, t, m, upl, wo_mode) -> np.ndarray:
    mask = np.ones(len(g), dtype=bool)
    if t is not None:
        mask &= (g["tkey"] == t).to_numpy()
    if m is not None:
        mask &= (g["mkey"] == m).to_numpy()
    if upl == "with":
        mask &= (g["upl"] == "Y").to_numpy()
    elif upl == "without":
        mask &= ((g["upl"] == "N") if wo_mode == "explicit" else (g["upl"] != "Y")).to_numpy()
    return mask


def _rate(g: pd.DataFrame, seg: np.ndarray, fam: str, v: dict, is_seg: bool) -> float:
    hf = g["hf"].to_numpy()
    under = ~np.isnan(hf) & (hf < HF_THRESHOLD_HOURS)
    batch = g["BATCHID"]

    keep = seg.copy()
    if v["creator"] == "all" or (v["creator"] == "segments" and is_seg):
        keep &= ~g["creator_bad"].to_numpy()

    if fam == "P":
        base = keep
    else:
        base = keep & g["audited"].to_numpy()
        if v["excl"]:
            base &= ~g["system"].to_numpy()

    if v["num"] == "all":
        num = batch[base].nunique()
    elif v["num"] == "under":
        num = batch[base & under].nunique()
    else:  # auditbase
        num = (batch[keep & g["abase_main"].to_numpy()].nunique()
               + batch[keep & g["abase_hold"].to_numpy()].nunique())

    if v["den"] == "batch":
        # SUMX(FILTER(VALUES(BATCHID), MAX(gap) < 0.25), MAX(gap)). A batch whose
        # max is BLANK passes DAX's "< 0.25" but adds 0, so dropping it is equal.
        mx = pd.Series(hf[base]).groupby(batch.to_numpy()[base], sort=False).max()
        den = float(mx[mx < HF_THRESHOLD_HOURS].sum())
    elif v["den"] == "row":
        dmask = base & under
        den = float((hf[dmask] * g["n"].to_numpy()[dmask]).sum())
    else:
        den = float(np.unique(hf[base & under]).sum())
    return float(num / den) if den > 0 else float("nan")


def _col_rate(g: pd.DataFrame, c: tuple, fam: str, v: dict, wo_mode) -> float:
    """Rate for one BENCH_COLUMNS entry. A column is a 'segment' if it filters on
    anything at all; only the family-wide overall rate isn't."""
    is_seg = any(x is not None for x in c[2:5])
    return _rate(g, _seg_mask(g, c[2], c[3], c[4], wo_mode), fam, v, is_seg)


def _variant_name(v: dict, fam: str) -> str:
    out = f"num={v['num']}, den={v['den']}, creator-filter={v['creator']}"
    if fam == "A":
        out += ", excl-SYSTEM-auditor" if v["excl"] else ", incl-SYSTEM-auditor"
    return out


def read_benchmark_sheet(path: Path):
    """Returns (workbook, sheet, header->column index, {Period: row index})."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    headers = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if isinstance(h, str):
            headers[" ".join(h.split()).lower()] = c
    rows = {}
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if hasattr(d, "year"):
            rows[pd.Period(year=d.year, month=d.month, freq="M")] = r
    return wb, ws, headers, rows


def _col_of(headers: dict, name: str):
    return headers.get(" ".join(name.split()).lower())


def run_benchmarks(args) -> int:
    print("=" * 74)
    print("KPI BENCHMARKS  (current month row)")
    print("=" * 74)
    xlsx = args.kpi
    print(f"  workbook: {xlsx}\n")
    if not xlsx.exists():
        print("  ERROR: workbook not found.\n", file=sys.stderr)
        return 1

    start = time.perf_counter()
    data, current = load_bench_data()
    if data is None or current is None:
        print("  ERROR: no Parquet for Current Raw Data -- nothing to compute.\n", file=sys.stderr)
        return 1
    months = sorted(data["__month"].unique())
    print(f"  loaded {len(data):,} rows, months {months[0]} .. {months[-1]} "
          f"({len(months)} months), current month {current}")

    # Diagnostics -- how the raw values mapped. Unmapped types are ignored.
    print("\n  DEPOSITTRANTYPE values -> key used:")
    tk = _norm_key(data["DEPOSITTRANTYPE"])
    for raw, n in data["DEPOSITTRANTYPE"].value_counts(dropna=False).head(12).items():
        k = tk[data["DEPOSITTRANTYPE"] == raw].iloc[0] if not pd.isna(raw) else None
        key = _TRANTYPE_KEYS.get(k, "(unmapped)") if isinstance(k, str) else "(blank)"
        print(f"    {str(raw)[:28]:28} {n:>10,}  -> {key}")
    print("  UPLOADABLEUSED values:  " + ",  ".join(
        f"{str(k)!s}={v:,}" for k, v in data["UPLOADABLEUSED"].value_counts(dropna=False).head(8).items()))
    print("  MODULE values:          " + ",  ".join(
        f"{str(k)!s}={v:,}" for k, v in data["MODULE"].value_counts(dropna=False).head(6).items()))

    frames = build_bench_frames(data)
    del data
    print(f"\n  gaps derived in {time.perf_counter() - start:.1f}s")

    wb, ws, headers, sheet_rows = read_benchmark_sheet(xlsx)
    missing = [c[0] for c in BENCH_COLUMNS if _col_of(headers, c[0]) is None]
    if missing:
        print(f"  WARNING: sheet has no column for: {', '.join(missing)}")

    # ---- calibration rows: historical rows whose window we mostly have ----
    avail = set(months)
    if args.calib_months:
        wanted = {pd.Period(m.strip(), freq="M") for m in args.calib_months.split(",")}
    else:
        wanted = None
    calib = []
    for per, r in sorted(sheet_rows.items()):
        if per >= current:
            continue
        win = _window(per)
        cov = sum(m in avail for m in win) / len(win)
        if (wanted is not None and per in wanted) or (wanted is None and cov >= BENCH_MIN_CALIB_COVERAGE):
            calib.append((per, r, win, cov))
    if not calib:
        print("  ERROR: no historical row has enough data to calibrate against.\n", file=sys.stderr)
        return 1
    print("  calibrating against: " + ",  ".join(f"{p} ({c:.0%} of window)" for p, _, _, c in calib))

    windows = {p: {f: frames[f][frames[f]["__month"].isin(w)] for f in frames} for p, _, w, _ in calib}

    chosen, fit_ok = {}, {}
    for fam in ("P", "A"):
        cols = [c for c in BENCH_COLUMNS if c[1] == fam and _col_of(headers, c[0])]
        # targets
        targets = {}
        for per, r, _, _ in calib:
            for c in cols:
                val = ws.cell(r, _col_of(headers, c[0])).value
                if isinstance(val, (int, float)):
                    targets[(per, c[0])] = float(val)

        results = []
        for v in BENCH_VARIANTS[fam]:
            errs = []
            for per, _, _, _ in calib:
                g = windows[per][fam]
                for c in cols:
                    if c[4] == "without":
                        continue  # scored separately below
                    tgt = targets.get((per, c[0]))
                    if tgt is None:
                        continue
                    got = _col_rate(g, c, fam, v, None)
                    if not np.isnan(got):
                        errs.append(abs(got - tgt) / tgt)
            results.append((np.mean(errs) if errs else np.inf, v))
        results.sort(key=lambda x: x[0])
        best_err, best = results[0]
        level_best_err, level_best = best_err, best
        pinned = BENCH_PINNED_VARIANTS.get(fam)
        if pinned is not None:
            for err, v in results:
                if v == pinned:
                    best_err, best = err, v
                    break

        # W/O mode, judged on the W/O columns only, with the chosen variant.
        wo_best, wo_err = WO_MODES[0], np.inf
        for mode in WO_MODES:
            errs = []
            for per, _, _, _ in calib:
                g = windows[per][fam]
                for c in cols:
                    tgt = targets.get((per, c[0]))
                    if c[4] != "without" or tgt is None:
                        continue
                    got = _col_rate(g, c, fam, best, mode)
                    if not np.isnan(got):
                        errs.append(abs(got - tgt) / tgt)
            e = np.mean(errs) if errs else np.inf
            if e < wo_err:
                wo_best, wo_err = mode, e

        chosen[fam] = (best, wo_best)
        fit_ok[fam] = best_err <= BENCH_FIT_TOLERANCE
        source = "from the report's DAX" if pinned is not None else "best fit"
        print(f"\n  {FAMILY_LABEL[fam]} -- using ({source}): {_variant_name(best, fam)}, W/O={wo_best}")
        print(f"    mean error {best_err:.2%}  (W/O columns {wo_err:.2%})  "
              f"-> {'MATCHES' if fit_ok[fam] else 'DOES NOT MATCH'} (tolerance {BENCH_FIT_TOLERANCE:.0%})"
              + ("  [chain mode: informational only]" if BENCH_STRATEGY == "chain" else ""))
        others = [(e, v) for e, v in results if v != best]
        for err, v in others[:3] if not args.bench_verbose else others:
            print(f"      alt       {err:7.2%}  {_variant_name(v, fam)}")

        # per-column detail on the best-covered calibration row
        per, _, _, cov = max(calib, key=lambda x: (x[3], x[0]))
        g = windows[per][fam]
        print(f"    per column vs {per} row ({cov:.0%} of window):")
        for c in cols:
            tgt = targets.get((per, c[0]))
            got = _col_rate(g, c, fam, best, wo_best)
            if tgt is None:
                print(f"      {c[0][:44]:44} sheet   (blank)  computed {got:8.2f}")
            else:
                print(f"      {c[0][:44]:44} sheet {tgt:8.2f}  computed {got:8.2f}  "
                      f"{(got - tgt) / tgt:+7.1%}")
        if len(calib) > 1:
            print("    mean error by row:  " + ",  ".join(
                f"{p}: " + f"{np.mean([abs(_col_rate(windows[p][fam], c, fam, best, wo_best) - targets[(p, c[0])]) / targets[(p, c[0])] for c in cols if (p, c[0]) in targets]):.1%}"
                for p, _, _, _ in calib))

    # ---- compute the current month ----
    cur_win = _window(current)
    have = [m for m in cur_win if m in avail]
    print(f"\n  current row {current}  window {_range_label(current)}  "
          f"({len(have)} of {len(cur_win)} months present)")
    if len(have) < len(cur_win):
        print(f"    missing: {', '.join(str(m) for m in cur_win if m not in avail)}")

    prev = current - 1
    prev_row = sheet_rows.get(prev)
    prev_win = _window(prev)
    chain = BENCH_STRATEGY == "chain" and prev_row is not None
    if BENCH_STRATEGY == "chain" and prev_row is None:
        print(f"  note: no {prev} row in the sheet to link from -- using raw computed values")
    if chain:
        prev_have = sum(m in avail for m in prev_win)
        print(f"  linking from the {prev} row  (its window: {prev_have} of {len(prev_win)} months present)")

    new_vals = {}   # name -> (family, raw computed, linked or None, sheet prev)
    for fam in ("P", "A"):
        g = frames[fam][frames[fam]["__month"].isin(cur_win)]
        gp = frames[fam][frames[fam]["__month"].isin(prev_win)]
        best, wo_mode = chosen[fam]
        for c in BENCH_COLUMNS:
            col = _col_of(headers, c[0])
            if c[1] != fam or col is None:
                continue
            raw = _col_rate(g, c, fam, best, wo_mode)
            sheet_prev = ws.cell(prev_row, col).value if prev_row else None
            linked = None
            if (chain and c[0] not in BENCH_RAW_COLUMNS
                    and isinstance(sheet_prev, (int, float)) and not np.isnan(raw)):
                comp_prev = _col_rate(gp, c, fam, best, wo_mode)
                if comp_prev and not np.isnan(comp_prev):
                    linked = sheet_prev * raw / comp_prev
            new_vals[c[0]] = (fam, raw, linked, sheet_prev)

    if chain:
        writable = {"P", "A"}
    else:
        writable = {f for f in fit_ok if fit_ok[f] or args.force_benchmarks}
    row = sheet_rows.get(current)
    is_new = row is None
    if is_new:
        row = ws.max_row + 1

    def _f(x):
        return f"{x:9.2f}" if isinstance(x, (int, float)) and not np.isnan(x) else f"{'--':>9}"

    print(f"\n  {'column':44} {str(prev) + ' sheet':>13} {'old':>9} {'raw':>9} {'NEW':>9} {'vs prev':>8}")
    print(f"  {'-' * 44} {'-' * 13} {'-' * 9} {'-' * 9} {'-' * 9} {'-' * 8}")
    final, jumps = {}, []
    for name, (fam, raw, linked, sheet_prev) in new_vals.items():
        col = _col_of(headers, name)
        old = None if is_new else ws.cell(row, col).value
        val = linked if linked is not None else raw
        val = None if val is None or np.isnan(val) else round(val, 2)
        final[name] = (fam, val)
        vs = (f"{(val - sheet_prev) / sheet_prev:+8.1%}"
              if val is not None and isinstance(sheet_prev, (int, float)) and sheet_prev else f"{'':>8}")
        flag = ""
        if fam not in writable:
            flag = "  (not written -- fit too poor)"
        elif name in BENCH_RAW_COLUMNS:
            flag = "  (raw by design -- corrected definition)"
        elif chain and linked is None:
            flag = "  (no link -- raw value used)"
        if (val is not None and isinstance(sheet_prev, (int, float)) and sheet_prev
                and name not in BENCH_RAW_COLUMNS
                and abs(val - sheet_prev) / sheet_prev > BENCH_JUMP_WARN):
            flag += f"  <-- moved >{BENCH_JUMP_WARN:.0%}, check"
            jumps.append(name)
        print(f"  {name[:44]:44} {_f(sheet_prev):>13} {_f(old)} {_f(raw)} {_f(val)} {vs}{flag}")
    if jumps:
        print(f"\n  WARNING: {len(jumps)} column(s) moved more than {BENCH_JUMP_WARN:.0%} from last month.")
        print("  Small segments (few batches) swing easily; check before trusting the chart.")
    if chain:
        print("\n  raw = computed from the data;  NEW = previous sheet value moved by the computed")
        print("  month-over-month change (BENCH_STRATEGY = \"chain\").  NEW is what gets written.")

    if args.dry_run:
        print("\n  --dry-run: workbook not modified.\n")
        return 0
    if not writable:
        print("\n  Neither family matched the historical rows closely enough, so the workbook\n"
              "  was NOT modified. Paste the output above back for another calibration round,\n"
              "  or re-run with --force-benchmarks to write the best-fit values anyway.\n")
        return 1

    # ---- write ----
    from copy import copy
    import shutil
    if is_new:
        src = row - 1
        for c in range(1, ws.max_column + 1):
            s, d = ws.cell(src, c), ws.cell(row, c)
            d.font, d.fill, d.border = copy(s.font), copy(s.fill), copy(s.border)
            d.alignment, d.number_format = copy(s.alignment), s.number_format
        ws.cell(row, 1).value = datetime.datetime(current.year, current.month, 1)
        ws.cell(row, 2).value = _range_label(current)
    for name, (fam, val) in final.items():
        if fam in writable and val is not None:
            ws.cell(row, _col_of(headers, name)).value = val

    KPI_BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    backup = KPI_BACKUP_DIR / f"KPI Benchmarks {datetime.datetime.now():%Y-%m-%d %H%M%S}.xlsx"
    shutil.copy2(xlsx, backup)
    try:
        wb.save(xlsx)
    except PermissionError:
        print("\n  ERROR: couldn't save -- KPI Benchmarks.xlsx is open somewhere. Close it in\n"
              "  Excel and run again. (Nothing was changed; a backup was still taken.)\n",
              file=sys.stderr)
        return 1
    print(f"\n  {'added' if is_new else 'updated'} row {current} ({_range_label(current)}); "
          f"backup: {backup.name}\n")
    return 0



def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--only", choices=[p.key for p in PROFILES] + ["controlsheet"],
                    help="run just one folder (or just the control sheet step)")
    ap.add_argument("--no-controlsheet", action="store_true", help="skip the control sheet step")
    ap.add_argument("--force", action="store_true", help="rebuild every file, ignoring timestamps")
    ap.add_argument("--audit", action="store_true", help="compare CSV schemas, then exit")
    ap.add_argument("--compression", default="snappy", choices=["snappy", "zstd", "gzip", "none"])
    ap.add_argument(
        "--keep-orphans", action="store_true", help="don't delete parquet whose CSV is gone"
    )
    ap.add_argument("--holidays", type=Path, default=HOLIDAYS_CSV, help="holiday CSV path")
    ap.add_argument("--kpi", type=Path, default=KPI_XLSX, help="KPI Benchmarks workbook path")
    ap.add_argument("--no-benchmarks", action="store_true", help="skip the KPI Benchmarks step")
    ap.add_argument("--benchmarks-only", action="store_true", help="skip the Parquet conversion")
    ap.add_argument("--dry-run", action="store_true", help="compute benchmarks but don't write")
    ap.add_argument("--force-benchmarks", action="store_true",
                    help="write benchmark values even if calibration doesn't match")
    ap.add_argument("--calib-months", default=None,
                    help="calibrate only on these rows, e.g. 2026-06,2026-05")
    ap.add_argument("--bench-verbose", action="store_true", help="list every calibration variant")
    args = ap.parse_args()

    profiles = [p for p in PROFILES if args.only is None or p.key == args.only]

    rc = 0
    if not args.benchmarks_only:
        for prof in profiles:
            rc |= run_profile(prof, args)
        if not (args.audit or args.no_controlsheet) and args.only in (None, "controlsheet"):
            try:
                rc |= run_control_sheet(args)
            except Exception as exc:
                import traceback
                traceback.print_exc()
                print(f"  Control sheet step FAILED: {type(exc).__name__}: {exc}\n", file=sys.stderr)
                rc |= 1

    if not (args.audit or args.no_benchmarks or args.only == "controlsheet"):
        try:
            rc |= run_benchmarks(args)
        except Exception as exc:
            import traceback
            traceback.print_exc()
            print(f"  KPI Benchmarks step FAILED: {type(exc).__name__}: {exc}\n", file=sys.stderr)
            rc |= 1

    print("All folders processed." if rc == 0 else "Finished with problems -- see above.")
    return rc


if __name__ == "__main__":
    sys.exit(main())
